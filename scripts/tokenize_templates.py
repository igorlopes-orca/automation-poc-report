"""Apply bullet tokens and strip baked-in hyperlinks from PT/ES templates.

Idempotent. Re-run after swapping in a fresh export of either deck. Each
run writes a one-time `.pptx.bak` backup.

Per slide entry:

- ``bullets``: per-language ordered list of bullet texts. Optional —
  some slides build their bullet block dynamically at fill time
  (e.g. slide 8 compliance frameworks); for those, omit ``bullets`` and
  only the link-stripping pass runs.

The ``BULLET_START_IDX`` for each slide lives in ``src.slide_layout``
so the runtime filler stays in sync with the template prep step.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402

from src.pptx_filler import rewrite_bullet_block  # noqa: E402
from src.slide_layout import BULLET_START_IDX  # noqa: E402

TEMPLATES = ROOT / "templates"

SLIDE_CONFIG: dict[int, dict] = {
    5: {
        "bullets": {
            "pt": [
                "{{metrics.vuln.exposed_crit_high_sensitive}} vulnerabilidades críticas ou altas em assets expostos com dados sensíveis",
                "{{metrics.vuln.exposed_crit_high}} vulnerabilidades críticas ou altas em assets expostos",
                "{{metrics.vuln.critical_with_fix}} vulnerabilidades críticas com fix disponível",
                "{{metrics.vuln.high_with_fix}} vulnerabilidades altas com fix disponível",
                "Total {{metrics.vuln.total}} vulnerabilidades",
            ],
            "es": [
                "{{metrics.vuln.exposed_crit_high_sensitive}} vulnerabilidades críticas o altas en activos expuestos con datos sensibles",
                "{{metrics.vuln.exposed_crit_high}} vulnerabilidades críticas o altas en activos expuestos",
                "{{metrics.vuln.critical_with_fix}} vulnerabilidades críticas con parche disponible",
                "{{metrics.vuln.high_with_fix}} vulnerabilidades altas con parche disponible",
                "Total {{metrics.vuln.total}} vulnerabilidades",
            ],
        },
    },
    6: {
        "bullets": {
            "pt": [
                "{{metrics.assets.malware}} ativos com malware",
                "{{metrics.assets.critical_exposed}} ativos com risco crítico expostos ao mundo",
                "{{metrics.assets.critical}} ativos com risco crítico",
                "{{metrics.assets.high_exposed}} ativos com risco alto expostos ao mundo",
                "{{metrics.assets.high}} ativos com risco alto",
            ],
            "es": [
                "{{metrics.assets.malware}} activos con malware",
                "{{metrics.assets.critical_exposed}} activos con riesgo crítico expuestos al mundo",
                "{{metrics.assets.critical}} activos con riesgo crítico",
                "{{metrics.assets.high_exposed}} activos con riesgo alto expuestos al mundo",
                "{{metrics.assets.high}} activos con riesgo alto",
            ],
        },
    },
    7: {
        # P5 is the "Top Findings:" header; bullets begin at P6. Only the
        # first two bullets are wired; the third (unused-permissions %)
        # will be added once its query lands.
        "bullets": {
            "pt": [
                "{{metrics.iam.permissive_identities}} identidades com acesso permissivo",
                "{{metrics.iam.attack_path_identities}} identidades que fazem parte de um caminho de ataque",
            ],
            "es": [
                "{{metrics.iam.permissive_identities}} identidades con acceso permisivo",
                "{{metrics.iam.attack_path_identities}} identidades que forman parte de un camino de ataque",
            ],
        },
    },
    8: {
        # Bullets built at fill time from the compliance API: 1 total
        # line + up to 5 framework lines. Tokenize only strips
        # hyperlinks; ``rewrite_bullet_block`` runs in fill_template.
    },
    9: {
        "bullets": {
            "pt": [
                "{{metrics.findings.crit_high_exposed_sensitive}} alertas críticos e altos em assets expostos com dados sensíveis",
                "{{metrics.findings.dspm_crit_high}} alertas críticos e altos de DSPM",
                "{{metrics.findings.validated_sensitive_data}} dados sensíveis validados pela Orca",
                "{{metrics.findings.high_impact_attack_paths}} attack paths de alto impacto",
                "{{metrics.findings.malware_crit_high}} alertas críticos e altos de malware",
            ],
            "es": [
                "{{metrics.findings.crit_high_exposed_sensitive}} alertas críticos y altos en activos expuestos con datos sensibles",
                "{{metrics.findings.dspm_crit_high}} alertas críticos y altos de DSPM",
                "{{metrics.findings.validated_sensitive_data}} datos sensibles validados por Orca",
                "{{metrics.findings.high_impact_attack_paths}} attack paths de alto impacto",
                "{{metrics.findings.malware_crit_high}} alertas críticos y altos de malware",
            ],
        },
    },
    10: {
        "bullets": {
            "pt": [
                "{{metrics.appsec.critical_alerts}} alertas críticos identificados em {{metrics.appsec.repos_with_critical}} repositórios",
                "{{metrics.appsec.high_alerts}} alertas altos identificados em {{metrics.appsec.repos_with_high}} repositórios",
                "{{metrics.appsec.sensitive_crit_high_alerts}} alertas de dados sensíveis críticos e altos identificados em {{metrics.appsec.repos_with_sensitive_crit_high}} repositórios",
                "{{metrics.appsec.repos_with_deployed_assets}} repositórios com recursos de infraestrutura mapeados com ambiente de produção",
            ],
            "es": [
                "{{metrics.appsec.critical_alerts}} alertas críticas identificadas en {{metrics.appsec.repos_with_critical}} repositorios",
                "{{metrics.appsec.high_alerts}} alertas altas identificadas en {{metrics.appsec.repos_with_high}} repositorios",
                "{{metrics.appsec.sensitive_crit_high_alerts}} alertas de datos sensibles críticas y altas identificadas en {{metrics.appsec.repos_with_sensitive_crit_high}} repositorios",
                "{{metrics.appsec.repos_with_deployed_assets}} repositorios con recursos de infraestructura asignados a un entorno de producción",
            ],
        },
    },
}


def _strip_hyperlinks(slide) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.hyperlink.address is not None:
                    run.hyperlink.address = None


def _apply_to_template(template_path: Path, lang: str) -> None:
    prs = Presentation(str(template_path))
    for slide_num, cfg in SLIDE_CONFIG.items():
        slide = prs.slides[slide_num - 1]
        _strip_hyperlinks(slide)
        if "bullets" in cfg:
            rewrite_bullet_block(slide, BULLET_START_IDX[slide_num], cfg["bullets"][lang])
    prs.save(str(template_path))


def main() -> None:
    for lang in ("pt", "es"):
        path = TEMPLATES / f"poc_report_{lang}.pptx"
        if not path.exists():
            raise FileNotFoundError(path)
        backup = path.with_suffix(".pptx.bak")
        if not backup.exists():
            shutil.copy2(path, backup)
            print(f"Backup: {backup}")
        _apply_to_template(path, lang)
        print(f"Tokenized: {path}")


if __name__ == "__main__":
    main()
