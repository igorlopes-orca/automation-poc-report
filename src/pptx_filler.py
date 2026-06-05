from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor

from src.slide_layout import BULLET_START_IDX
from src.slide_styling import SLIDE_STYLING, SlideStyling, apply_slide_styling

TOKEN_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}")
BLACK = RGBColor(0x00, 0x00, 0x00)


def _resolve(path: str, data: dict) -> Any:
    """Walk dotted path (e.g. 'metrics.vuln.total') through nested dicts."""
    node: Any = data
    for part in path.split("."):
        if not isinstance(node, dict) or part not in node:
            raise KeyError(f"Token path not found in data: {path}")
        node = node[part]
    return node


def _substitute_paragraph(paragraph, data: dict) -> None:
    """Substitute ``{{...}}`` tokens in a paragraph.

    PowerPoint often splits a token like ``{{x.y.z}}`` across three runs
    (``{{``, ``x.y.z``, ``}}``), so per-run substitution misses it. We
    reconstruct the full paragraph text, substitute across the whole
    string, then collapse into the first run — preserving that run's
    font (colour, size, bold). Callers should tokenize paragraphs that
    are visually uniform; that matches how our variable fields are
    authored in the template.
    """
    if not paragraph.runs:
        return
    full = "".join(run.text for run in paragraph.runs)
    if "{{" not in full:
        return
    new_full = TOKEN_RE.sub(lambda m: str(_resolve(m.group(1), data)), full)
    if new_full == full:
        return

    paragraph.runs[0].text = new_full
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)


def _fill_text_frames(slide, data: dict) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            _substitute_paragraph(paragraph, data)


def _first_text_frame(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            return shape.text_frame
    raise RuntimeError("No text frame on slide")


def _set_paragraph_text(paragraph, new_text: str) -> None:
    """Overwrite a paragraph's text using a single run.

    Forces the run to plain black and strips any hyperlink — bullets are
    a neutral base; the styler owns colour/link decisions.
    """
    runs = paragraph.runs
    if not runs:
        raise RuntimeError("Paragraph has no runs; cannot rewrite without losing formatting.")
    runs[0].text = new_text
    runs[0].font.color.rgb = BLACK
    if runs[0].hyperlink.address is not None:
        runs[0].hyperlink.address = None
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def rewrite_bullet_block(slide, bullet_start_idx: int, bullets: list[str]) -> None:
    """Resize the bullet section to exactly ``len(bullets)`` paragraphs
    (cloning or pruning) and overwrite each.

    Used by the template tokenizer (slides 5/6/7) and by ``fill_template``
    for slides whose bullet text is built at runtime (slide 8).
    """
    tf = _first_text_frame(slide)
    paragraphs = list(tf.paragraphs)

    current = len(paragraphs) - bullet_start_idx
    target = len(bullets)

    if current < target:
        template_p = paragraphs[-1]._p
        for _ in range(target - current):
            clone = deepcopy(template_p)
            template_p.addnext(clone)
            template_p = clone
    elif current > target:
        for paragraph in paragraphs[bullet_start_idx + target :]:
            paragraph._p.getparent().remove(paragraph._p)

    paragraphs = list(tf.paragraphs)
    for i, text in enumerate(bullets):
        _set_paragraph_text(paragraphs[bullet_start_idx + i], text)


def fill_template(
    template_path: Path,
    output_path: Path,
    data: dict,
    bullet_blocks: dict[int, list[str]] | None = None,
    styling: dict[int, SlideStyling] | None = None,
) -> Path:
    """Open template, apply text tokens, dynamic bullet blocks, then
    slide styling, and save.

    Args:
        template_path: .pptx to use as the starting point.
        output_path: where to save the filled deck.
        data: nested dict; tokens like ``{{metrics.vuln.total}}`` resolve
            via dotted-path lookup.
        bullet_blocks: per-slide bullet lists for slides whose content is
            built at runtime (e.g. slide 8 compliance frameworks). Keyed
            by slide number (1-based). The ``bullet_start_idx`` is read
            from ``slide_layout.BULLET_START_IDX``.
        styling: per-slide colour/hyperlink rules (see ``slide_styling``).
            Defaults to ``SLIDE_STYLING`` from that module.
    """
    prs = Presentation(str(template_path))
    for slide in prs.slides:
        _fill_text_frames(slide, data)

    for slide_num, bullets in (bullet_blocks or {}).items():
        slide = prs.slides[slide_num - 1]
        rewrite_bullet_block(slide, BULLET_START_IDX[slide_num], bullets)

    apply_slide_styling(prs, styling if styling is not None else SLIDE_STYLING)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path
